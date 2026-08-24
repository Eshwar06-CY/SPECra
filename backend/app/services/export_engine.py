"""
UniHack 252 Static Expected Output Headers Specification and Dynamic Mapping Engine.

Preserves exact header naming, spelling, case, and ordering as defined in
'Unihack_ Expected Output - Delivery Format.csv'.
"""
from typing import Any, Dict, List, Optional
import io
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from app.models.product import Product, ProductAttribute, Evidence, ValidationResult


# Exact 252 UniHack static output headers in precise order
UNIHACK_STATIC_HEADERS: List[str] = [
    "MFR URL", "Ref URL 1", "Ref URL 2", "Ref URL 3", "Ref URL 4", "Ref URL 5",
    "PART_NUMBER", "Dept", "Class", "Fine", "SKU - MY_PART_NUMBER", "Mfg_Part_Num",
    "Part_Desc", "E1_Brand", "Unilog_Brand", "DIB_Brand", "Part_Manuf",
    "MANUFACTURER_NAME", "BRAND_NAME", "TRADE_NAME", "MANUFACTURER_PART_NUMBER",
    "ALTERNATE_PART_NUMBER", "Classpath", "MOBILE_DESC", "INVOICE_DESC",
    "SHORT_DESC", "LONG_DESC1", "RETAIL_DESC", "MARKETING_DESCRIPTION",
    "ITEM_FEATURES_1", "ITEM_FEATURES_2", "ITEM_FEATURES_3", "ITEM_FEATURES_4",
    "ITEM_FEATURES_5", "ITEM_FEATURES_6", "ITEM_FEATURES_7", "ITEM_FEATURES_8",
    "ITEM_FEATURES_9", "ITEM_FEATURES_10", "ITEM_FEATURES_11", "ITEM_FEATURES_12",
    "ITEM_FEATURES_13", "ITEM_FEATURES_14", "ITEM_FEATURES_15", "ITEM_FEATURES_16",
    "ITEM_FEATURES_17", "ITEM_FEATURES_18", "ITEM_FEATURES_19", "ITEM_FEATURES_20",
    "With", "Standard/Approvals", "Prop 65", "Application", "Includes", "Product Name",
    # 50 dynamic attribute slots (LABEL, VALUE, UOM) = 150 headers
    "ATTRIBUTE_LABEL 1", "ATTRIBUTE_VALUE 1", "ATTRIBUTE_UOM 1",
    "ATTRIBUTE_LABEL 2", "ATTRIBUTE_VALUE 2", "ATTRIBUTE_UOM 2",
    "ATTRIBUTE_LABEL 3", "ATTRIBUTE_VALUE 3", "ATTRIBUTE_UOM 3",
    "ATTRIBUTE_LABEL 4", "ATTRIBUTE_VALUE 4", "ATTRIBUTE_UOM 4",
    "ATTRIBUTE_LABEL 5", "ATTRIBUTE_VALUE 5", "ATTRIBUTE_UOM 5",
    "ATTRIBUTE_LABEL 6", "ATTRIBUTE_VALUE 6", "ATTRIBUTE_UOM 6",
    "ATTRIBUTE_LABEL 7", "ATTRIBUTE_VALUE 7", "ATTRIBUTE_UOM 7",
    "ATTRIBUTE_LABEL 8", "ATTRIBUTE_VALUE 8", "ATTRIBUTE_UOM 8",
    "ATTRIBUTE_LABEL 9", "ATTRIBUTE_VALUE 9", "ATTRIBUTE_UOM 9",
    "ATTRIBUTE_LABEL 10", "ATTRIBUTE_VALUE 10", "ATTRIBUTE_UOM 10",
    "ATTRIBUTE_LABEL 11", "ATTRIBUTE_VALUE 11", "ATTRIBUTE_UOM 11",
    "ATTRIBUTE_LABEL 12", "ATTRIBUTE_VALUE 12", "ATTRIBUTE_UOM 12",
    "ATTRIBUTE_LABEL 13", "ATTRIBUTE_VALUE 13", "ATTRIBUTE_UOM 13",
    "ATTRIBUTE_LABEL 14", "ATTRIBUTE_VALUE 14", "ATTRIBUTE_UOM 14",
    "ATTRIBUTE_LABEL 15", "ATTRIBUTE_VALUE 15", "ATTRIBUTE_UOM 15",
    "ATTRIBUTE_LABEL 16", "ATTRIBUTE_VALUE 16", "ATTRIBUTE_UOM 16",
    "ATTRIBUTE_LABEL 17", "ATTRIBUTE_VALUE 17", "ATTRIBUTE_UOM 17",
    "ATTRIBUTE_LABEL 18", "ATTRIBUTE_VALUE 18", "ATTRIBUTE_UOM 18",
    "ATTRIBUTE_LABEL 19", "ATTRIBUTE_VALUE 19", "ATTRIBUTE_UOM 19",
    "ATTRIBUTE_LABEL 20", "ATTRIBUTE_VALUE 20", "ATTRIBUTE_UOM 20",
    "ATTRIBUTE_LABEL 21", "ATTRIBUTE_VALUE 21", "ATTRIBUTE_UOM 21",
    "ATTRIBUTE_LABEL 22", "ATTRIBUTE_VALUE 22", "ATTRIBUTE_UOM 22",
    "ATTRIBUTE_LABEL 23", "ATTRIBUTE_VALUE 23", "ATTRIBUTE_UOM 23",
    "ATTRIBUTE_LABEL 24", "ATTRIBUTE_VALUE 24", "ATTRIBUTE_UOM 24",
    "ATTRIBUTE_LABEL 25", "ATTRIBUTE_VALUE 25", "ATTRIBUTE_UOM 25",
    "ATTRIBUTE_LABEL 26", "ATTRIBUTE_VALUE 26", "ATTRIBUTE_UOM 26",
    "ATTRIBUTE_LABEL 27", "ATTRIBUTE_VALUE 27", "ATTRIBUTE_UOM 27",
    "ATTRIBUTE_LABEL 28", "ATTRIBUTE_VALUE 28", "ATTRIBUTE_UOM 28",
    "ATTRIBUTE_LABEL 29", "ATTRIBUTE_VALUE 29", "ATTRIBUTE_UOM 29",
    "ATTRIBUTE_LABEL 30", "ATTRIBUTE_VALUE 30", "ATTRIBUTE_UOM 30",
    "ATTRIBUTE_LABEL 31", "ATTRIBUTE_VALUE 31", "ATTRIBUTE_UOM 31",
    "ATTRIBUTE_LABEL 32", "ATTRIBUTE_VALUE 32", "ATTRIBUTE_UOM 32",
    "ATTRIBUTE_LABEL 33", "ATTRIBUTE_VALUE 33", "ATTRIBUTE_UOM 33",
    "ATTRIBUTE_LABEL 34", "ATTRIBUTE_VALUE 34", "ATTRIBUTE_UOM 34",
    "ATTRIBUTE_LABEL 35", "ATTRIBUTE_VALUE 35", "ATTRIBUTE_UOM 35",
    "ATTRIBUTE_LABEL 36", "ATTRIBUTE_VALUE 36", "ATTRIBUTE_UOM 36",
    "ATTRIBUTE_LABEL 37", "ATTRIBUTE_VALUE 37", "ATTRIBUTE_UOM 37",
    "ATTRIBUTE_LABEL 38", "ATTRIBUTE_VALUE 38", "ATTRIBUTE_UOM 38",
    "ATTRIBUTE_LABEL 39", "ATTRIBUTE_VALUE 39", "ATTRIBUTE_UOM 39",
    "ATTRIBUTE_LABEL 40", "ATTRIBUTE_VALUE 40", "ATTRIBUTE_UOM 40",
    "ATTRIBUTE_LABEL 41", "ATTRIBUTE_VALUE 41", "ATTRIBUTE_UOM 41",
    "ATTRIBUTE_LABEL 42", "ATTRIBUTE_VALUE 42", "ATTRIBUTE_UOM 42",
    "ATTRIBUTE_LABEL 43", "ATTRIBUTE_VALUE 43", "ATTRIBUTE_UOM 43",
    "ATTRIBUTE_LABEL 44", "ATTRIBUTE_VALUE 44", "ATTRIBUTE_UOM 44",
    "ATTRIBUTE_LABEL 45", "ATTRIBUTE_VALUE 45", "ATTRIBUTE_UOM 45",
    "ATTRIBUTE_LABEL 46", "ATTRIBUTE_VALUE 46", "ATTRIBUTE_UOM 46",
    "ATTRIBUTE_LABEL 47", "ATTRIBUTE_VALUE 47", "ATTRIBUTE_UOM 47",
    "ATTRIBUTE_LABEL 48", "ATTRIBUTE_VALUE 48", "ATTRIBUTE_UOM 48",
    "ATTRIBUTE_LABEL 49", "ATTRIBUTE_VALUE 49", "ATTRIBUTE_UOM 49",
    "ATTRIBUTE_LABEL 50", "ATTRIBUTE_VALUE 50", "ATTRIBUTE_UOM 50",
    # Commerce, Identifiers & Packaging
    "UPC", "EAN", "GTIN", "UNSPSC", "Warranty", "List Price",
    "Selling Qty", "Selling UOM", "Standard Packaging Information",
    # Physical Specifications & Dimensions
    "LENGTH", "LENGTH_UOM", "HEIGHT", "HEIGHT_UOM",
    "WIDTH", "WIDTH_UOM", "WEIGHT", "WEIGHT_UOM", "VOLUME", "VOLUME_UOM",
    # Digital Assets, Media & Compliance Manuals
    "Product Image", "Alternate Image 1", "Alternate Image 2", "Alternate Image 3", "Alternate Image 4",
    "SDS", "SDS_1", "Warranty Information", "Catalog", "Specification Sheet",
    "Instruction/Installation Manual", "Service Manual", "Owners/User Manual",
    "Line Drawing", "MTR", "RoHS", "Full Engineering Drawing", "Energy Star Guide",
    "Technical Bulletin", "Submittal", "Compatibility Chart", "Size Chart",
    "Product Label/Insert", "Video Link", "Video Link 1",
    "Country Of Origin", "Discontinued", "Actual Image (Yes/No)"
]


class UniHackOutputMapper:
    """
    Transforms database Product, ProductAttribute, Evidence, and Validation models
    into authoritative 252-column UniHack rows without inventing values.
    """

    @classmethod
    def map_product_to_row(
        cls,
        product: Product,
    ) -> Dict[str, str]:
        """
        Maps a single Product and its associated attributes/evidence into a 252-key dictionary.
        Leaves unpopulated columns as empty string ('').
        """
        row: Dict[str, str] = {h: "" for h in UNIHACK_STATIC_HEADERS}
        raw_data: Dict[str, Any] = product.raw_data or {}

        # 1. Map Raw Input passthrough columns
        for raw_k, raw_v in raw_data.items():
            if raw_v is not None:
                # Direct match
                if raw_k in row:
                    row[raw_k] = str(raw_v)
                # Case-insensitive match for raw keys
                for h in UNIHACK_STATIC_HEADERS[:20]:
                    if h.lower() == raw_k.lower() and not row[h]:
                        row[h] = str(raw_v)

        # 2. Master Product Identity & Classification
        if product.product_name:
            row["Product Name"] = str(product.product_name)
        if product.category:
            row["Class"] = str(product.category)

        # MPN / Part Numbers
        mpn = product.external_product_id or raw_data.get("Mfg_Part_Num") or raw_data.get("PART_NUMBER")
        if mpn:
            row["MANUFACTURER_PART_NUMBER"] = str(mpn)
            if not row["PART_NUMBER"]:
                row["PART_NUMBER"] = str(mpn)
            if not row["Mfg_Part_Num"]:
                row["Mfg_Part_Num"] = str(mpn)
            if not row["SKU - MY_PART_NUMBER"]:
                row["SKU - MY_PART_NUMBER"] = str(mpn)

        # Collect structured attributes and features
        identity_names = {"brand", "manufacturer", "product_type", "manufacturer_part_number", "part_number", "sku"}
        spec_attributes: List[ProductAttribute] = []
        feature_attributes: List[ProductAttribute] = []

        for attr in product.attributes:
            name_lower = attr.attribute_name.lower().strip()

            # Separate identity fields
            if name_lower == "brand":
                row["BRAND_NAME"] = str(attr.attribute_value or "")
            elif name_lower == "manufacturer":
                row["MANUFACTURER_NAME"] = str(attr.attribute_value or "")
                if not row["Part_Manuf"]:
                    row["Part_Manuf"] = str(attr.attribute_value or "")
            elif name_lower in ("manufacturer_part_number", "part_number") and not row["MANUFACTURER_PART_NUMBER"]:
                row["MANUFACTURER_PART_NUMBER"] = str(attr.attribute_value or "")
            elif name_lower == "product_type" and not row["Class"]:
                row["Class"] = str(attr.attribute_value or "")
            elif name_lower.startswith("feature_"):
                feature_attributes.append(attr)
            elif name_lower not in identity_names:
                spec_attributes.append(attr)

        # 3. Features Mapping (ITEM_FEATURES_1 to ITEM_FEATURES_20)
        for idx, feat in enumerate(feature_attributes[:20]):
            feat_header = f"ITEM_FEATURES_{idx + 1}"
            if feat_header in row:
                row[feat_header] = str(feat.attribute_value or "")

        # 4. Physical Dimensions Mapping
        for attr in spec_attributes:
            attr_name = attr.attribute_name.lower()
            val_str = str(attr.normalized_value or attr.attribute_value or "")
            unit_str = str(attr.unit or "")

            if attr_name == "length":
                row["LENGTH"] = val_str
                row["LENGTH_UOM"] = unit_str
            elif attr_name == "width":
                row["WIDTH"] = val_str
                row["WIDTH_UOM"] = unit_str
            elif attr_name == "height":
                row["HEIGHT"] = val_str
                row["HEIGHT_UOM"] = unit_str
            elif attr_name == "weight":
                row["WEIGHT"] = val_str
                row["WEIGHT_UOM"] = unit_str
            elif attr_name == "volume":
                row["VOLUME"] = val_str
                row["VOLUME_UOM"] = unit_str
            elif attr_name in ("pack_quantity", "quantity", "selling_quantity"):
                row["Selling Qty"] = val_str
                row["Selling UOM"] = unit_str or "pieces"

        # 5. Dynamic Attributes Slots Mapping (ATTRIBUTE_LABEL 1..50, ATTRIBUTE_VALUE 1..50, ATTRIBUTE_UOM 1..50)
        for idx, attr in enumerate(spec_attributes[:50]):
            slot_num = idx + 1
            label_header = f"ATTRIBUTE_LABEL {slot_num}"
            value_header = f"ATTRIBUTE_VALUE {slot_num}"
            uom_header = f"ATTRIBUTE_UOM {slot_num}"

            if label_header in row and not row[label_header]:
                row[label_header] = attr.attribute_name.replace("_", " ").title()
            if value_header in row and not row[value_header]:
                row[value_header] = str(attr.normalized_value or attr.attribute_value or "")
            if uom_header in row and not row[uom_header]:
                row[uom_header] = str(attr.unit or "")

        # 6. Apply Validated Deterministic Enrichments (if present on product)
        if hasattr(product, "enrichments") and product.enrichments:
            for enrich in product.enrichments:
                f_name = enrich.field_name
                if f_name in row and not row[f_name]:
                    row[f_name] = str(enrich.value or "")

        # 7. Deterministic Fallback for Unanalyzed / Bulk Catalog Rows
        # Enables high fill rate without running expensive per-item AI on standard fields
        part_desc = str(raw_data.get("Part_Desc") or "").strip()
        part_manuf = str(raw_data.get("Part_Manuf") or "").strip()
        mfg_num = str(raw_data.get("Mfg_Part_Num") or "").strip()

        from app.services.enrichment_engine import ProductEnrichmentEngine

        # Fallback Manufacturer Name
        if not row["MANUFACTURER_NAME"]:
            norm_m = ProductEnrichmentEngine.normalize_manufacturer_name(part_manuf)
            if norm_m:
                row["MANUFACTURER_NAME"] = norm_m

        # Fallback Brand Name
        if not row["BRAND_NAME"]:
            brand_val, _, _ = ProductEnrichmentEngine.extract_deterministic_brand(raw_data, part_desc)
            if brand_val:
                row["BRAND_NAME"] = brand_val

        # Fallback Product Name
        if not row["Product Name"]:
            norm_title = ProductEnrichmentEngine.normalize_product_title(part_desc, mfg_num)
            if norm_title:
                row["Product Name"] = norm_title

        # Fallback Class
        if not row["Class"]:
            tax_class = ProductEnrichmentEngine.classify_product_taxonomy(part_desc)
            if tax_class:
                row["Class"] = tax_class

        # Helper to find next empty attribute slot
        def _get_next_empty_attr_slot(r: Dict[str, str]) -> Optional[int]:
            for s in range(1, 51):
                if not r.get(f"ATTRIBUTE_LABEL {s}") and not r.get(f"ATTRIBUTE_VALUE {s}"):
                    return s
            return None

        # Fallback Dimensions & Technical Attributes
        dim_info = ProductEnrichmentEngine.extract_deterministic_dimensions(part_desc)
        if dim_info:
            d_type = dim_info.get("type")
            if d_type == "rectangular":
                if not row["WIDTH"] and "WIDTH" in dim_info:
                    row["WIDTH"] = dim_info["WIDTH"]
                    row["WIDTH_UOM"] = dim_info.get("WIDTH_UOM", "in")
                if not row["LENGTH"] and "LENGTH" in dim_info:
                    row["LENGTH"] = dim_info["LENGTH"]
                    row["LENGTH_UOM"] = dim_info.get("LENGTH_UOM", "in")
            elif d_type in ("diameter", "three_axis"):
                # Circular tools: populate Diameter, Thickness, Arbor in dynamic attribute slots
                if "diameter" in dim_info:
                    # Check if already present
                    has_diam = any(row.get(f"ATTRIBUTE_LABEL {s}") == "Diameter" for s in range(1, 51))
                    if not has_diam:
                        slot = _get_next_empty_attr_slot(row)
                        if slot:
                            row[f"ATTRIBUTE_LABEL {slot}"] = "Diameter"
                            row[f"ATTRIBUTE_VALUE {slot}"] = dim_info["diameter"]
                            row[f"ATTRIBUTE_UOM {slot}"] = dim_info.get("diameter_uom", "in")

                if "thickness" in dim_info:
                    has_thick = any(row.get(f"ATTRIBUTE_LABEL {s}") == "Thickness" for s in range(1, 51))
                    if not has_thick:
                        slot = _get_next_empty_attr_slot(row)
                        if slot:
                            row[f"ATTRIBUTE_LABEL {slot}"] = "Thickness"
                            row[f"ATTRIBUTE_VALUE {slot}"] = dim_info["thickness"]
                            row[f"ATTRIBUTE_UOM {slot}"] = dim_info.get("thickness_uom", "in")

                if "arbor" in dim_info:
                    has_arbor = any(row.get(f"ATTRIBUTE_LABEL {s}") == "Arbor Hole Size" for s in range(1, 51))
                    if not has_arbor:
                        slot = _get_next_empty_attr_slot(row)
                        if slot:
                            row[f"ATTRIBUTE_LABEL {slot}"] = "Arbor Hole Size"
                            row[f"ATTRIBUTE_VALUE {slot}"] = dim_info["arbor"]
                            row[f"ATTRIBUTE_UOM {slot}"] = dim_info.get("arbor_uom", "in")

        # Fallback Packaging
        if not row["Selling Qty"]:
            pkg_info = ProductEnrichmentEngine.extract_deterministic_packaging(part_desc)
            if pkg_info:
                row["Selling Qty"] = pkg_info["qty"]
                row["Selling UOM"] = pkg_info["uom"]
                if not row["Standard Packaging Information"]:
                    row["Standard Packaging Information"] = pkg_info["description"]

        # Fallback Grit in Dynamic Attribute Slot
        has_grit = any(row.get(f"ATTRIBUTE_LABEL {s}") == "Grit Size" for s in range(1, 51))
        if not has_grit:
            grit_info = ProductEnrichmentEngine.extract_deterministic_grit(part_desc)
            if grit_info:
                slot = _get_next_empty_attr_slot(row)
                if slot:
                    row[f"ATTRIBUTE_LABEL {slot}"] = "Grit Size"
                    row[f"ATTRIBUTE_VALUE {slot}"] = grit_info["value"]
                    row[f"ATTRIBUTE_UOM {slot}"] = "Grit"

        return row


def sanitize_spreadsheet_cell(val: Any) -> str:
    """
    Sanitizes string values to prevent CSV / Excel formula injection (CSV Injection).
    If a cell begins with dangerous formula characters ('=', '+', '-', '@', '\t', '\r'),
    it prepends an apostrophe single-quote (') so spreadsheet software treats it as pure text.
    Numbers and clean strings remain unaffected.
    """
    if val is None:
        return ""
    str_val = str(val)
    if not str_val:
        return ""

    # Dangerous formula prefixes in Excel/LibreOffice/Google Sheets
    DANGEROUS_PREFIXES = ("=", "+", "-", "@", "\t", "\r")

    # Don't sanitize standard pure numbers (e.g. -5 or +12) if they are purely numeric
    trimmed = str_val.strip()
    if trimmed.startswith(("-", "+")):
        try:
            float(trimmed)
            return str_val
        except ValueError:
            pass

    if str_val.startswith(DANGEROUS_PREFIXES):
        return f"'{str_val}"

    return str_val


class UniHackExportEngine:
    """
    Generates downloadable CSV and XLSX files conforming strictly to UniHack 252 static headers
    with built-in formula injection mitigation.
    """

    @classmethod
    def generate_csv_bytes(cls, rows: List[Dict[str, str]]) -> bytes:
        """Serializes mapped product rows into UTF-8-SIG encoded CSV bytes."""
        output = io.StringIO()
        writer = csv.writer(output, lineterminator="\n")

        # Write exact 252 header line
        writer.writerow(UNIHACK_STATIC_HEADERS)

        # Write product rows with formula sanitization
        for row in rows:
            row_values = [sanitize_spreadsheet_cell(row.get(h, "")) for h in UNIHACK_STATIC_HEADERS]
            writer.writerow(row_values)

        return output.getvalue().encode("utf-8-sig")

    @classmethod
    def generate_xlsx_bytes(cls, rows: List[Dict[str, str]]) -> bytes:
        """Serializes mapped product rows into styled Microsoft Excel (XLSX) bytes."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "UniHack Delivery Output"

        # Header styling
        header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        data_font = Font(name="Segoe UI", size=9)
        data_alignment = Alignment(horizontal="left", vertical="center")

        thin_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0"),
        )

        # Write header row
        ws.append(UNIHACK_STATIC_HEADERS)
        header_row = ws[1]
        for cell in header_row:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment

        ws.row_dimensions[1].height = 28

        # Write data rows with formula sanitization
        for r_idx, row in enumerate(rows, start=2):
            row_values = [sanitize_spreadsheet_cell(row.get(h, "")) for h in UNIHACK_STATIC_HEADERS]
            ws.append(row_values)
            for c_idx in range(1, len(UNIHACK_STATIC_HEADERS) + 1):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.font = data_font
                cell.alignment = data_alignment
                cell.border = thin_border
            ws.row_dimensions[r_idx].height = 20

        # Freeze top header row
        ws.freeze_panes = "A2"

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()
