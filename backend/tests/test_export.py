"""
Unit tests for the UniHack 252-Column Output Mapping and Export Engine.
Runs 100% offline without calling any LLM API.
"""
import unittest
import uuid
import io
import csv
import openpyxl

from app.models.product import Product, ProductAttribute, Evidence, ProcessingJob, ValidationResult
from app.services.export_engine import UNIHACK_STATIC_HEADERS, UniHackOutputMapper, UniHackExportEngine
from app.services.export_service import ExportService
from app.core.database import SessionLocal


class TestExportEngine(unittest.TestCase):
    """
    Comprehensive offline test suite for the Export Engine.
    """

    def setUp(self):
        self.db = SessionLocal()

    def tearDown(self):
        self.db.rollback()
        self.db.close()

    def test_exact_252_headers_and_order(self):
        """Test 1 & 2: Preserves exactly 252 static headers and strict ordering"""
        self.assertEqual(len(UNIHACK_STATIC_HEADERS), 252)
        self.assertEqual(UNIHACK_STATIC_HEADERS[0], "MFR URL")
        self.assertEqual(UNIHACK_STATIC_HEADERS[6], "PART_NUMBER")
        self.assertEqual(UNIHACK_STATIC_HEADERS[11], "Mfg_Part_Num")
        self.assertEqual(UNIHACK_STATIC_HEADERS[12], "Part_Desc")
        self.assertEqual(UNIHACK_STATIC_HEADERS[17], "MANUFACTURER_NAME")
        self.assertEqual(UNIHACK_STATIC_HEADERS[18], "BRAND_NAME")
        self.assertEqual(UNIHACK_STATIC_HEADERS[20], "MANUFACTURER_PART_NUMBER")
        self.assertEqual(UNIHACK_STATIC_HEADERS[29], "ITEM_FEATURES_1")
        self.assertEqual(UNIHACK_STATIC_HEADERS[48], "ITEM_FEATURES_20")
        self.assertEqual(UNIHACK_STATIC_HEADERS[54], "Product Name")
        self.assertEqual(UNIHACK_STATIC_HEADERS[55], "ATTRIBUTE_LABEL 1")
        self.assertEqual(UNIHACK_STATIC_HEADERS[56], "ATTRIBUTE_VALUE 1")
        self.assertEqual(UNIHACK_STATIC_HEADERS[57], "ATTRIBUTE_UOM 1")
        self.assertEqual(UNIHACK_STATIC_HEADERS[204], "ATTRIBUTE_UOM 50")
        self.assertEqual(UNIHACK_STATIC_HEADERS[214], "LENGTH")
        self.assertEqual(UNIHACK_STATIC_HEADERS[218], "WIDTH")
        self.assertEqual(UNIHACK_STATIC_HEADERS[251], "Actual Image (Yes/No)")

    def test_single_product_mapping(self):
        """Test 3: Single product maps canonical identity, attributes, features, and dimensions"""
        product = Product(
            id=uuid.uuid4(),
            external_product_id="DCB518ASTS06G",
            product_name="Diablo 1/2 in x 18 in Sanding Belt 6pc",
            category="Sanding Belt",
            raw_data={
                "Mfg_Part_Num": "DCB518ASTS06G",
                "Part_Desc": "Diablo Sanding Belt 6pc",
                "Part_Manuf": "Freud Inc (2435)",
            },
        )
        attr_brand = ProductAttribute(product_id=product.id, attribute_name="brand", attribute_value="Diablo")
        attr_manuf = ProductAttribute(product_id=product.id, attribute_name="manufacturer", attribute_value="Freud Inc (2435)")
        attr_width = ProductAttribute(product_id=product.id, attribute_name="width", attribute_value="1/2 in", normalized_value="0.5", unit="in")
        attr_len = ProductAttribute(product_id=product.id, attribute_name="length", attribute_value="18 in", normalized_value="18", unit="in")
        attr_qty = ProductAttribute(product_id=product.id, attribute_name="pack_quantity", attribute_value="6 pieces", normalized_value="6", unit="pieces")
        feat1 = ProductAttribute(product_id=product.id, attribute_name="feature_durability", attribute_value="Heavy duty cloth backing")

        product.attributes = [attr_brand, attr_manuf, attr_width, attr_len, attr_qty, feat1]

        row = UniHackOutputMapper.map_product_to_row(product)

        # Verify key fields mapped
        self.assertEqual(len(row), 252)
        self.assertEqual(row["Product Name"], "Diablo 1/2 in x 18 in Sanding Belt 6pc")
        self.assertEqual(row["Class"], "Sanding Belt")
        self.assertEqual(row["BRAND_NAME"], "Diablo")
        self.assertEqual(row["MANUFACTURER_NAME"], "Freud Inc (2435)")
        self.assertEqual(row["MANUFACTURER_PART_NUMBER"], "DCB518ASTS06G")
        self.assertEqual(row["Mfg_Part_Num"], "DCB518ASTS06G")
        self.assertEqual(row["WIDTH"], "0.5")
        self.assertEqual(row["WIDTH_UOM"], "in")
        self.assertEqual(row["LENGTH"], "18")
        self.assertEqual(row["LENGTH_UOM"], "in")
        self.assertEqual(row["ITEM_FEATURES_1"], "Heavy duty cloth backing")
        self.assertEqual(row["Selling Qty"], "6")
        self.assertEqual(row["Selling UOM"], "pieces")

        # Dynamic attribute slot
        self.assertEqual(row["ATTRIBUTE_LABEL 1"], "Width")
        self.assertEqual(row["ATTRIBUTE_VALUE 1"], "0.5")
        self.assertEqual(row["ATTRIBUTE_UOM 1"], "in")

    def test_csv_generation(self):
        """Test 5: CSV generation returns valid bytes with 252 columns and correct row count"""
        rows = [
            {"Product Name": "Product A", "BRAND_NAME": "Brand A", "MANUFACTURER_PART_NUMBER": "PN-A"},
            {"Product Name": "Product B", "BRAND_NAME": "Brand B", "MANUFACTURER_PART_NUMBER": "PN-B"},
        ]
        csv_bytes = UniHackExportEngine.generate_csv_bytes(rows)
        self.assertIsInstance(csv_bytes, bytes)
        
        # Decode and parse CSV
        content = csv_bytes.decode("utf-8-sig")
        reader = list(csv.reader(io.StringIO(content)))
        self.assertEqual(len(reader), 3)  # Header + 2 rows
        self.assertEqual(len(reader[0]), 252)
        self.assertEqual(len(reader[1]), 252)
        self.assertEqual(len(reader[2]), 252)

    def test_xlsx_generation(self):
        """Test 6: XLSX generation returns valid openpyxl workbook bytes"""
        rows = [
            {"Product Name": "Product A", "BRAND_NAME": "Brand A", "MANUFACTURER_PART_NUMBER": "PN-A"},
        ]
        xlsx_bytes = UniHackExportEngine.generate_xlsx_bytes(rows)
        self.assertIsInstance(xlsx_bytes, bytes)

        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
        ws = wb.active
        self.assertEqual(ws.max_column, 252)
        self.assertEqual(ws.max_row, 2)  # Header + 1 row
        self.assertEqual(ws.cell(row=1, column=1).value, "MFR URL")

    def test_empty_values_handled_cleanly(self):
        """Test 7: Unpopulated fields remain empty strings without crashing or hallucinating"""
        empty_prod = Product(id=uuid.uuid4(), product_name="Minimal Product", raw_data={})
        empty_prod.attributes = []

        row = UniHackOutputMapper.map_product_to_row(empty_prod)
        self.assertEqual(row["Product Name"], "Minimal Product")
        self.assertEqual(row["BRAND_NAME"], "")
        self.assertEqual(row["MFR URL"], "")
        self.assertEqual(row["Prop 65"], "")
        self.assertEqual(len(row), 252)

    def test_arbitrary_input_columns_passthrough(self):
        """Test 11: Arbitrary input column combinations are passed through cleanly"""
        custom_prod = Product(
            id=uuid.uuid4(),
            product_name="Custom Valve",
            raw_data={
                "MFR URL": "https://example.com/item",
                "E1_Brand": "CustomBrand",
                "Custom_Col_XYZ": "Ignored or Custom",
            },
        )
        custom_prod.attributes = []
        row = UniHackOutputMapper.map_product_to_row(custom_prod)
        self.assertEqual(row["MFR URL"], "https://example.com/item")
        self.assertEqual(row["E1_Brand"], "CustomBrand")

    def test_no_unexpected_columns(self):
        """Test 13: Export row contains strictly the 252 expected keys"""
        prod = Product(id=uuid.uuid4(), product_name="Test Item", raw_data={})
        prod.attributes = []
        row = UniHackOutputMapper.map_product_to_row(prod)
        self.assertEqual(set(row.keys()), set(UNIHACK_STATIC_HEADERS))

    def test_preview_export_integration(self):
        """Test 14: ExportService.preview_export returns correct 252-header structure without crashing"""
        job = ProcessingJob(
            id=uuid.uuid4(),
            filename="sample_test_catalog.csv",
            file_type="csv",
            status="completed",
            total_records=1,
            processed_records=1,
            failed_records=0,
        )
        self.db.add(job)
        self.db.flush()

        product = Product(
            id=uuid.uuid4(),
            job_id=job.id,
            product_name="Test Industrial Abrasive",
            category="Abrasives",
            raw_data={"PART_NUMBER": "3M-TEST-01", "Part_Desc": "Sanding Disc 50/Pk"},
        )
        self.db.add(product)
        self.db.commit()

        # Call preview_export
        preview = ExportService.preview_export(db=self.db, job_id=job.id, limit=1)

        self.assertIn("job_id", preview)
        self.assertEqual(preview["job_id"], str(job.id))
        self.assertEqual(preview["filename"], "sample_test_catalog.csv")
        self.assertEqual(preview["total_products"], 1)
        self.assertEqual(preview["total_headers"], 252)
        self.assertIn("summary", preview)
        self.assertIn("sample_rows", preview)
        self.assertEqual(len(preview["sample_rows"]), 1)
        self.assertEqual(len(preview["sample_rows"][0]), 252)
        self.assertEqual(preview["sample_rows"][0]["Product Name"], "Test Industrial Abrasive")

    def test_15_deterministic_manufacturer_normalization(self):
        """Test 15: Normalizes 'Freud Inc (2435)' -> 'Freud Inc' without losing company names"""
        prod = Product(
            id=uuid.uuid4(),
            raw_data={"Mfg_Part_Num": "DCB518", "Part_Manuf": "Freud Inc (2435)"}
        )
        row = UniHackOutputMapper.map_product_to_row(prod)
        self.assertEqual(row["MANUFACTURER_NAME"], "Freud Inc")
        self.assertEqual(row["Part_Manuf"], "Freud Inc (2435)")

    def test_16_deterministic_packaging_extraction(self):
        """Test 16: Extracts '50 Disc/Box' -> 50 / disc / 50 pieces per box"""
        prod = Product(
            id=uuid.uuid4(),
            raw_data={"Mfg_Part_Num": "3M775L", "Part_Desc": "3M 775L Stikit Film P80 - Cubitron II 50 Disc/Box"}
        )
        row = UniHackOutputMapper.map_product_to_row(prod)
        self.assertEqual(row["Selling Qty"], "50")
        self.assertEqual(row["Selling UOM"], "disc")
        self.assertEqual(row["Standard Packaging Information"], "50 pieces per box")

    def test_17_deterministic_dimensions_extraction(self):
        """Test 17: Extracts '1/2\"x18\"' -> width 0.5 in / length 18 in"""
        prod = Product(
            id=uuid.uuid4(),
            raw_data={"Mfg_Part_Num": "DCB518", "Part_Desc": "Diablo 1/2\"x18\" - Sanding Belt 6pc"}
        )
        row = UniHackOutputMapper.map_product_to_row(prod)
        self.assertEqual(row["WIDTH"], "0.5")
        self.assertEqual(row["WIDTH_UOM"], "in")
        self.assertEqual(row["LENGTH"], "18")
        self.assertEqual(row["LENGTH_UOM"], "in")
        self.assertEqual(row["Selling Qty"], "6")
        self.assertEqual(row["Class"], "Sanding Belt")

    def test_18_deterministic_brand_and_grit(self):
        """Test 18: Extracts Brand '3M' and Grit 'P80' into dynamic slot 1"""
        prod = Product(
            id=uuid.uuid4(),
            raw_data={"Mfg_Part_Num": "3M775L", "Part_Desc": "3M 775L Stikit Film P80 - Cubitron II 50 Disc/Box"}
        )
        row = UniHackOutputMapper.map_product_to_row(prod)
        self.assertEqual(row["BRAND_NAME"], "3M")
        self.assertEqual(row["ATTRIBUTE_LABEL 1"], "Grit Size")
        self.assertEqual(row["ATTRIBUTE_VALUE 1"], "80")
        self.assertEqual(row["ATTRIBUTE_UOM 1"], "Grit")

    def test_20_precision_three_axis_cut_off_disc(self):
        """TEST 4: '49-94-0013 Milw 5\"x.045\"x7/8\" Metal Cut Off Disc' -> Diameter=5, Thickness=0.045, Arbor=7/8, Width=blank"""
        prod = Product(
            id=uuid.uuid4(),
            raw_data={"Mfg_Part_Num": "49-94-0013", "Part_Desc": "49-94-0013 Milw 5\"x.045\"x7/8\" Metal Cut Off Disc"}
        )
        row = UniHackOutputMapper.map_product_to_row(prod)
        self.assertEqual(row["WIDTH"], "")
        self.assertEqual(row["LENGTH"], "")
        self.assertEqual(row["ATTRIBUTE_LABEL 1"], "Diameter")
        self.assertEqual(row["ATTRIBUTE_VALUE 1"], "5")
        self.assertEqual(row["ATTRIBUTE_UOM 1"], "in")
        self.assertEqual(row["ATTRIBUTE_LABEL 2"], "Thickness")
        self.assertEqual(row["ATTRIBUTE_VALUE 2"], "0.045")
        self.assertEqual(row["ATTRIBUTE_UOM 2"], "in")
        self.assertEqual(row["ATTRIBUTE_LABEL 3"], "Arbor Hole Size")
        self.assertEqual(row["ATTRIBUTE_VALUE 3"], "7/8")
        self.assertEqual(row["ATTRIBUTE_UOM 3"], "in")

    def test_21_precision_three_axis_4inch_cutoff(self):
        """TEST 5: '4\"x.040\"x5/8\" Cut Off Disc' -> Diameter=4, Thickness=0.04, Arbor=5/8"""
        prod = Product(
            id=uuid.uuid4(),
            raw_data={"Mfg_Part_Num": "49-94-0001", "Part_Desc": "4\"x.040\"x5/8\" Cut Off Disc"}
        )
        row = UniHackOutputMapper.map_product_to_row(prod)
        self.assertEqual(row["WIDTH"], "")
        self.assertEqual(row["LENGTH"], "")
        self.assertEqual(row["ATTRIBUTE_LABEL 1"], "Diameter")
        self.assertEqual(row["ATTRIBUTE_VALUE 1"], "4")
        self.assertEqual(row["ATTRIBUTE_LABEL 2"], "Thickness")
        self.assertEqual(row["ATTRIBUTE_VALUE 2"], "0.04")
        self.assertEqual(row["ATTRIBUTE_LABEL 3"], "Arbor Hole Size")
        self.assertEqual(row["ATTRIBUTE_VALUE 3"], "5/8")

    def test_22_precision_circular_disc_diameter(self):
        """TEST 3: 'HIOLIT 5\" P80' -> Diameter=5, Grit Size=80, WIDTH=blank, LENGTH=blank"""
        prod = Product(
            id=uuid.uuid4(),
            raw_data={"Mfg_Part_Num": "5B-332-080", "Part_Desc": "5B-332-080 HIOLIT 5\" P80"}
        )
        row = UniHackOutputMapper.map_product_to_row(prod)
        self.assertEqual(row["WIDTH"], "")
        self.assertEqual(row["LENGTH"], "")
        # Should have Grit Size and Diameter in dynamic slots
        labels = [row.get(f"ATTRIBUTE_LABEL {s}") for s in range(1, 5)]
        self.assertIn("Grit Size", labels)
        self.assertIn("Diameter", labels)
        self.assertEqual(row["ATTRIBUTE_VALUE 1"], "5")
        self.assertEqual(row["ATTRIBUTE_VALUE 2"], "80")

    def test_23_precision_rectangular_belt(self):
        """TEST 2: '2.75x30 Sanding Belt' -> WIDTH=2.75, LENGTH=30"""
        prod = Product(
            id=uuid.uuid4(),
            raw_data={"Mfg_Part_Num": "9A-570-240", "Part_Desc": "9A-570-240 Abranet 2.75x30 Sanding Belt"}
        )
        row = UniHackOutputMapper.map_product_to_row(prod)
        self.assertEqual(row["WIDTH"], "2.75")
        self.assertEqual(row["WIDTH_UOM"], "in")
        self.assertEqual(row["LENGTH"], "30")
        self.assertEqual(row["LENGTH_UOM"], "in")

    def test_24_precision_sheets_packaging(self):
        """TEST 6: '9A-129-120 Abranet 3x4 - 50 Sheets/Box' -> Selling Qty=50, Selling UOM=sheets, Packaging='50 sheets per box'"""
        prod = Product(
            id=uuid.uuid4(),
            raw_data={"Mfg_Part_Num": "9A-129-120", "Part_Desc": "9A-129-120 Abranet 3x4 - 50 Sheets/Box"}
        )
        row = UniHackOutputMapper.map_product_to_row(prod)
        self.assertEqual(row["Selling Qty"], "50")
        self.assertEqual(row["Selling UOM"], "sheets")
        self.assertEqual(row["Standard Packaging Information"], "50 sheets per box")
        self.assertEqual(row["WIDTH"], "3")
        self.assertEqual(row["LENGTH"], "4")


if __name__ == "__main__":
    unittest.main()
